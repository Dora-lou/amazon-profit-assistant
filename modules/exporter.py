from __future__ import annotations

from datetime import date
from io import BytesIO
from re import sub

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from modules.calculations import (
    breakeven_price,
    breakeven_tacos,
    fixed_cost,
    max_tacos_for_margin,
    product_from_row,
    profit_breakdown,
    target_margin_price,
)


EXPORT_COLUMNS = [
    ("产品名称", "text"),
    ("ASIN", "text"),
    ("FNSKU", "text"),
    ("固定成本($)", "usd"),
    ("平台佣金(%)", "pct"),
    ("仓储率(%)", "pct"),
    ("退货率(%)", "pct"),
    ("划线价 / 当前售价($)", "usd"),
    ("预估销量(件)", "num0"),
    ("当前TACOS", "pct"),
    ("目标TACOS", "pct"),
    ("目标净利率", "pct"),
    ("单件净利($)", "usd"),
    ("单件利润率(%)", "pct"),
    ("保本TACOS", "pct"),
    ("目标利润TACOS上限", "pct"),
    ("绝对保本售价($)", "usd"),
    ("目标利润最低售价($)", "usd"),
    ("月预估销售额($)", "usd"),
    ("月预估利润($)", "usd"),
]


def safe_filename(name: str) -> str:
    cleaned = sub(r'[\\/:*?"<>|]+', "_", name).strip()
    return cleaned or "Amazon产品"


def export_filename(products: list[dict]) -> str:
    today = date.today().isoformat()
    if len(products) == 1:
        product = product_from_row(products[0])
        return f"{safe_filename(product.name)}_利润测算_{today}.xlsx"
    return f"Amazon利润测算_{today}.xlsx"


def row_values(row: dict) -> list:
    product = product_from_row(row)
    breakdown = profit_breakdown(product)
    return [
        product.name,
        product.asin,
        product.fnsku,
        fixed_cost(product),
        product.commission_rate,
        product.storage_rate,
        product.return_rate,
        product.price,
        product.daily_sales,
        product.tacos,
        product.target_tacos,
        product.target_margin,
        breakdown["net_profit"],
        breakdown["net_margin"],
        breakeven_tacos(product),
        max_tacos_for_margin(product, product.target_margin),
        breakeven_price(product),
        target_margin_price(product),
        breakdown["monthly_revenue"],
        breakdown["monthly_profit"],
    ]


def build_profit_workbook(products: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "利润测算"

    group_fill = {
        "cost": PatternFill("solid", fgColor="D9EAF7"),
        "ops": PatternFill("solid", fgColor="E2F0D9"),
        "decision": PatternFill("solid", fgColor="FCE4D6"),
    }
    header_fill = PatternFill("solid", fgColor="F3F4F6")
    positive_font = Font(color="008000")
    negative_font = Font(color="C00000")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws.merge_cells(start_row=1, start_column=8, end_row=1, end_column=14)
    ws.merge_cells(start_row=1, start_column=15, end_row=1, end_column=20)
    group_titles = [(1, "1. 成本", "cost"), (8, "2. 日常经营 / 利润", "ops"), (15, "3. 经营决策", "decision")]
    for col, title, key in group_titles:
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = group_fill[key]
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col_index, (header, _) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, product in enumerate(products, start=3):
        for col_index, value in enumerate(row_values(product), start=1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            number_type = EXPORT_COLUMNS[col_index - 1][1]
            if number_type == "cny":
                cell.number_format = '"￥"#,##0.00'
            elif number_type == "usd":
                cell.number_format = '"$"#,##0.00'
            elif number_type == "pct":
                cell.number_format = "0.0%"
            elif number_type == "num0":
                cell.number_format = "#,##0"
            elif number_type == "num":
                cell.number_format = "#,##0.00"

            if EXPORT_COLUMNS[col_index - 1][0] in {"单件净利($)", "月预估利润($)"} and isinstance(value, (int, float)):
                cell.font = negative_font if value < 0 else positive_font

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            text = "" if cell.value is None else str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 0), min(max(len(text) + 2, 10), 24))
    for col_index, width in widths.items():
        ws.column_dimensions[get_column_letter(col_index)].width = width

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
